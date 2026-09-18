from app.rag.ingestion.table_aware_parser import split_regions
from app.schemas.ingestion import HeaderSource, RegionType
from tests.fixtures.documents import (
    make_docx_bytes_with_table,
    make_pdf_bytes_with_table,
)


def _html(body: str) -> bytes:
    return f"<html><body>{body}</body></html>".encode()


def test_html_table_with_th_row_is_explicit_header() -> None:
    content = _html(
        "<table>"
        "<tr><th>Name</th><th>Score</th></tr>"
        "<tr><td>Alice</td><td>90</td></tr>"
        "<tr><td>Bob</td><td>85</td></tr>"
        "</table>"
    )

    [region] = split_regions(content, "handbook.html", "html")

    assert region.table is not None
    assert region.table.header_source == HeaderSource.EXPLICIT
    assert region.table.header_confidence == 1.0
    assert region.table.header_row == ["Name", "Score"]
    assert region.table.data_rows == [["Alice", "90"], ["Bob", "85"]]


def test_html_table_without_th_is_missing_header() -> None:
    content = _html(
        "<table>"
        "<tr><td>Name</td><td>Score</td></tr>"
        "<tr><td>Alice</td><td>90</td></tr>"
        "</table>"
    )

    [region] = split_regions(content, "handbook.html", "html")

    assert region.table is not None
    assert region.table.header_source == HeaderSource.MISSING
    assert region.table.header_row is None
    assert region.table.header_confidence == 0.0
    assert region.table.data_rows == [["Name", "Score"], ["Alice", "90"]]


def test_docx_table_without_tblheader_attribute_is_missing() -> None:
    """`add_table` fixture never sets `<w:tblHeader/>` - regular python-docx
    tables are not automatically flagged as having a header row just
    because they have a plausible-looking first row."""

    content = make_docx_bytes_with_table("Intro paragraph.", "Outro paragraph.")

    regions = split_regions(content, "handbook.docx", "docx")
    table_region = next(r for r in regions if r.region_type == RegionType.TABLE)

    assert table_region.table is not None
    assert table_region.table.header_source == HeaderSource.MISSING
    assert table_region.table.header_row is None
    assert table_region.table.data_rows == [
        ["Name", "Score"],
        ["Alice", "90"],
        ["Bob", "85"],
    ]


def test_docx_table_with_tblheader_attribute_is_explicit_header() -> None:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document()
    document.add_paragraph("Intro.")
    table = document.add_table(rows=3, cols=2)
    for row, cells in zip(
        table.rows, [["Name", "Score"], ["Alice", "90"], ["Bob", "85"]], strict=True
    ):
        for cell, value in zip(row.cells, cells, strict=True):
            cell.text = value
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tr_pr.append(tr_pr.makeelement(qn("w:tblHeader"), {}))
    import io

    buffer = io.BytesIO()
    document.save(buffer)

    regions = split_regions(buffer.getvalue(), "handbook.docx", "docx")
    table_region = next(r for r in regions if r.region_type == RegionType.TABLE)

    assert table_region.table is not None
    assert table_region.table.header_source == HeaderSource.EXPLICIT
    assert table_region.table.header_confidence == 1.0
    assert table_region.table.header_row == ["Name", "Score"]
    assert table_region.table.data_rows == [["Alice", "90"], ["Bob", "85"]]


def test_pdf_table_is_always_inferred_header() -> None:
    content = make_pdf_bytes_with_table("Intro paragraph.", "Outro paragraph.")

    regions = split_regions(content, "handbook.pdf", "pdf")
    table_region = next(r for r in regions if r.region_type == RegionType.TABLE)

    assert table_region.table is not None
    assert table_region.table.header_source == HeaderSource.INFERRED
    assert table_region.table.header_confidence == 0.6
    assert table_region.table.header_row is not None
    assert any("Alice" in row for row in table_region.table.data_rows)
