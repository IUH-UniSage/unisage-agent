import io

import openpyxl
import pymupdf
from docx import Document


def make_xlsx_bytes(header: list[str], rows: list[list[object]]) -> bytes:
    """Build a minimal in-memory `.xlsx` with one header row and data rows."""

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_pdf_bytes(text: str) -> bytes:
    """Build a minimal in-memory single-page PDF containing `text`."""

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    data: bytes = document.tobytes()
    document.close()
    return data


def make_pdf_bytes_with_table(intro_text: str, outro_text: str) -> bytes:
    """Build a PDF with an intro paragraph, a 2-column/3-row grid table, then an outro paragraph."""

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), intro_text)

    x0, y0, cell_w, cell_h, rows, cols = 72, 120, 80, 20, 3, 2
    for row in range(rows + 1):
        page.draw_line((x0, y0 + row * cell_h), (x0 + cols * cell_w, y0 + row * cell_h))
    for col in range(cols + 1):
        page.draw_line((x0 + col * cell_w, y0), (x0 + col * cell_w, y0 + rows * cell_h))

    table_cells = [["Name", "Score"], ["Alice", "90"], ["Bob", "85"]]
    for row, cells in enumerate(table_cells):
        for col, value in enumerate(cells):
            page.insert_text((x0 + col * cell_w + 5, y0 + row * cell_h + 14), value, fontsize=9)

    page.insert_text((72, 260), outro_text)

    data: bytes = document.tobytes()
    document.close()
    return data


def make_docx_bytes(text: str) -> bytes:
    """Build a minimal in-memory `.docx` containing one paragraph."""

    document = Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_docx_bytes_with_table(intro_text: str, outro_text: str) -> bytes:
    """Build a `.docx` with an intro paragraph, a real Word table, then an outro paragraph."""

    document = Document()
    document.add_paragraph(intro_text)
    table = document.add_table(rows=3, cols=2)
    for row, cells in zip(
        table.rows, [["Name", "Score"], ["Alice", "90"], ["Bob", "85"]], strict=True
    ):
        for cell, value in zip(row.cells, cells, strict=True):
            cell.text = value
    document.add_paragraph(outro_text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
